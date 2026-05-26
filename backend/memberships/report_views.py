import csv
import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.views import View

from .models import AuctionReport, AuctionTransaction, Donation
from .report_serializers import (
    AuctionTransactionDetailReportSerializer,
    AuctionTransactionReportDataSerializer,
)


EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def normalize_pdf_output(pdf_output):
    if isinstance(pdf_output, bytearray):
        return bytes(pdf_output)
    if isinstance(pdf_output, str):
        return pdf_output.encode("latin-1")
    return pdf_output


def fit_pdf_cell_text(pdf, value, width, style="", max_size=8, min_size=5):
    text = " ".join(str(value if value not in [None, ""] else "-").split())
    available_width = max(width - 2, 1)

    for font_size in range(max_size, min_size - 1, -1):
        pdf.set_font("Arial", style, size=font_size)
        if pdf.get_string_width(text) <= available_width:
            return text, font_size

    pdf.set_font("Arial", style, size=min_size)
    ellipsis = "..."
    if pdf.get_string_width(ellipsis) > available_width:
        return "", min_size

    trimmed = text
    while trimmed and pdf.get_string_width(f"{trimmed}{ellipsis}") > available_width:
        trimmed = trimmed[:-1]
    return f"{trimmed.rstrip()}{ellipsis}", min_size


def draw_pdf_table_header(pdf, columns):
    pdf.set_fill_color(139, 94, 52)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Arial", "B", size=8)
    draw_pdf_wrapped_row(pdf, columns, [label for label, _width, _max_length in columns], fill=True, align="C")
    pdf.set_text_color(0, 0, 0)


def draw_pdf_wrapped_row(pdf, columns, values, fill=False, align="L"):
    row_height = 7
    cell_padding = 1
    font_style = "B" if fill else ""
    fitted_values = [
        fit_pdf_cell_text(pdf, value, width, style=font_style)
        for (_label, width, _max_length), value in zip(columns, values)
    ]

    page_break_trigger = getattr(pdf, "page_break_trigger", 185)
    if pdf.get_y() + row_height > page_break_trigger:
        pdf.add_page()
        if not fill:
            draw_pdf_table_header(pdf, columns)

    x_start = pdf.get_x()
    y_start = pdf.get_y()

    for (_label, width, _max_length), fitted_value in zip(columns, fitted_values):
        x = pdf.get_x()
        y = pdf.get_y()
        pdf.rect(x, y, width, row_height, "DF" if fill else "")
        text, font_size = fitted_value
        pdf.set_font("Arial", font_style, size=font_size)
        pdf.set_xy(x + cell_padding, y + 2)
        pdf.cell(width - (cell_padding * 2), 3, text, border=0, align=align)
        pdf.set_xy(x + width, y)

    pdf.set_xy(x_start, y_start + row_height)


def draw_pdf_table_row(pdf, columns, values):
    pdf.set_text_color(0, 0, 0)
    draw_pdf_wrapped_row(pdf, columns, values)


def make_json_ready(value):
    return json.loads(json.dumps(value, default=str))


def get_limit_offset(request):
    try:
        limit = int(request.GET.get("limit", 10))
    except (TypeError, ValueError):
        limit = 10
    try:
        offset = int(request.GET.get("offset", 0))
    except (TypeError, ValueError):
        offset = 0
    return max(limit, 1), max(offset, 0)


def money_decimal(value):
    return float(value or Decimal("0.00"))


def get_transaction_source_id(transaction):
    if transaction.member:
        return transaction.member.member_id
    if transaction.relative:
        return transaction.relative.non_member_id
    return "N/A"


def get_transaction_source_name(transaction):
    if transaction.member:
        return transaction.member.name
    if transaction.relative:
        return transaction.relative.name
    return "N/A"


def get_transaction_source_type(transaction):
    return "Member" if transaction.member else "Non-Member"


def get_donation_source_id(donation):
    if donation.member:
        return donation.member.member_id
    if donation.relative:
        return donation.relative.non_member_id
    return "N/A"


def get_filtered_donations(search=""):
    queryset = Donation.objects.select_related("member", "relative").all()
    if not search:
        return queryset
    return queryset.filter(
        Q(donor_type__icontains=search)
        | Q(member__member_id__icontains=search)
        | Q(member__name__icontains=search)
        | Q(relative__non_member_id__icontains=search)
        | Q(relative__name__icontains=search)
        | Q(donor_name__icontains=search)
        | Q(phone__icontains=search)
    )


def get_transaction_payment_date(transaction):
    receipt_date = transaction.receipt.receipt_date if transaction.receipt else None
    return receipt_date.strftime("%Y-%m-%d") if receipt_date else ""


def get_transaction_receipt_no(transaction):
    return transaction.receipt.receipt_no if transaction.receipt else "-"


def build_excel_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(output.getvalue(), content_type=EXCEL_MIME_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def apply_excel_table_style(worksheet, row_number, column_count, header_fill, header_font, border):
    for col in range(1, column_count + 1):
        cell = worksheet.cell(row=row_number, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = worksheet.parent._report_center_alignment


def apply_excel_body_style(worksheet, row_number, column_count, border):
    for col in range(1, column_count + 1):
        cell = worksheet.cell(row=row_number, column=col)
        cell.border = border
        cell.alignment = worksheet.parent._report_body_alignment


def create_payment_status_excel_workbook(data, status_label):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    summary = data["summary"]
    transactions = data["transactions"]

    workbook = openpyxl.Workbook()
    workbook._report_center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    workbook._report_body_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet = workbook.active
    worksheet.title = f"{status_label.title()} Report"
    worksheet.freeze_panes = "A19"

    header_fill = PatternFill(start_color="8B5E34", end_color="8B5E34", fill_type="solid")
    title_fill = PatternFill(start_color="342515", end_color="342515", fill_type="solid")
    section_fill = PatternFill(start_color="F1E6D8", end_color="F1E6D8", fill_type="solid")
    positive_fill = PatternFill(start_color="E7F5EC", end_color="E7F5EC", fill_type="solid")
    warning_fill = PatternFill(start_color="FCE8E1", end_color="FCE8E1", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(size=16, bold=True, color="FFFFFF")
    section_font = Font(size=12, bold=True, color="2B2116")
    label_font = Font(bold=True, color="2B2116")
    thin_side = Side(style="thin", color="D8C8B8")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    worksheet.merge_cells("A1:J1")
    title_cell = worksheet["A1"]
    title_cell.value = f"{status_label} AUCTION TRANSACTION REPORT"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = workbook._report_center_alignment
    worksheet.row_dimensions[1].height = 28

    worksheet["A2"] = "Generated Date"
    worksheet["A2"].font = label_font
    worksheet["B2"] = summary["report_date"].strftime("%d-%m-%Y %I:%M %p")
    worksheet["B2"].number_format = "dd-mm-yyyy h:mm AM/PM"

    worksheet.merge_cells("A4:J4")
    worksheet["A4"] = "Summary"
    worksheet["A4"].font = section_font
    worksheet["A4"].fill = section_fill

    summary_cards = [
        ("Report Status", summary["status"], None),
        (f"{status_label.title()} Transactions", summary["status_count"], None),
        (f"{status_label.title()} Amount", money_decimal(summary["status_amount"]), "money"),
        ("All Transactions", summary["total_transactions"], None),
        ("All Amount", money_decimal(summary["total_amount"]), "money"),
        (f"Other ({summary['other_label']}) Count", summary["other_count"], None),
        (f"Other ({summary['other_label']}) Amount", money_decimal(summary["other_amount"]), "money"),
        ("Paid %", summary["paid_percentage"] / 100, "percent"),
        ("Unpaid %", summary["unpaid_percentage"] / 100, "percent"),
    ]

    for index, (label, value, value_type) in enumerate(summary_cards):
        row = 5 + index // 3
        col = 1 + (index % 3) * 3
        label_cell = worksheet.cell(row=row, column=col)
        value_cell = worksheet.cell(row=row, column=col + 1)
        worksheet.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 2)
        label_cell.value = label
        value_cell.value = value
        label_cell.font = label_font
        label_cell.fill = section_fill
        if "Paid" in label:
            value_cell.fill = positive_fill
        elif "Unpaid" in label:
            value_cell.fill = warning_fill
        if value_type == "money":
            value_cell.number_format = '"Rs." #,##0.00'
        elif value_type == "percent":
            value_cell.number_format = "0.00%"
        for style_col in range(col, col + 3):
            worksheet.cell(row=row, column=style_col).border = border
            worksheet.cell(row=row, column=style_col).alignment = workbook._report_body_alignment

    worksheet.merge_cells("A10:D10")
    worksheet["A10"] = "Payment Status Breakdown"
    worksheet["A10"].font = section_font
    worksheet["A10"].fill = section_fill

    breakdown_header_row = 11
    breakdown_headers = ["Status", "Count", "Total Amount", "Percentage"]
    for col, header in enumerate(breakdown_headers, 1):
        worksheet.cell(row=breakdown_header_row, column=col).value = header
    apply_excel_table_style(worksheet, breakdown_header_row, len(breakdown_headers), header_fill, header_font, border)

    row = breakdown_header_row + 1
    for item in data["payment_status_breakdown"]:
        worksheet.cell(row=row, column=1).value = item["payment_status"]
        worksheet.cell(row=row, column=2).value = item["count"]
        worksheet.cell(row=row, column=3).value = money_decimal(item["total_amount"])
        worksheet.cell(row=row, column=3).number_format = '"Rs." #,##0.00'
        worksheet.cell(row=row, column=4).value = item["percentage"] / 100
        worksheet.cell(row=row, column=4).number_format = "0.00%"
        apply_excel_body_style(worksheet, row, len(breakdown_headers), border)
        row += 1

    transaction_title_row = row + 2
    worksheet.merge_cells(start_row=transaction_title_row, start_column=1, end_row=transaction_title_row, end_column=10)
    worksheet.cell(row=transaction_title_row, column=1).value = f"{status_label.title()} Transaction Details"
    worksheet.cell(row=transaction_title_row, column=1).font = section_font
    worksheet.cell(row=transaction_title_row, column=1).fill = section_fill

    transaction_header_row = transaction_title_row + 1
    transaction_headers = [
        "S.No",
        "Transaction ID",
        "Member/Non-Member ID",
        "Name",
        "Type",
        "Auction Item",
        "Token No",
        "Price",
        "Payment Status",
        "Payment Date",
    ]
    for col, header in enumerate(transaction_headers, 1):
        worksheet.cell(row=transaction_header_row, column=col).value = header
    apply_excel_table_style(worksheet, transaction_header_row, len(transaction_headers), header_fill, header_font, border)

    row = transaction_header_row + 1
    for serial_number, txn in enumerate(transactions, 1):
        values = [
            serial_number,
            txn.id,
            get_transaction_source_id(txn),
            get_transaction_source_name(txn),
            get_transaction_source_type(txn),
            txn.item.auction_item_name,
            txn.token_number,
            money_decimal(txn.price),
            txn.payment_status,
            get_transaction_payment_date(txn),
        ]
        for col, value in enumerate(values, 1):
            worksheet.cell(row=row, column=col).value = value
        worksheet.cell(row=row, column=7).number_format = '"Rs." #,##0.00'
        apply_excel_body_style(worksheet, row, len(transaction_headers), border)
        row += 1

    last_data_row = max(row - 1, transaction_header_row)
    worksheet.auto_filter.ref = f"A{transaction_header_row}:J{last_data_row}"

    for column_letter, width in {
        "A": 8,
        "B": 14,
        "C": 22,
        "D": 28,
        "E": 15,
        "F": 28,
        "G": 12,
        "H": 16,
        "I": 16,
        "J": 22,
    }.items():
        worksheet.column_dimensions[column_letter].width = width

    for row_number in range(1, last_data_row + 1):
        worksheet.row_dimensions[row_number].height = 22

    worksheet.sheet_view.showGridLines = False
    return workbook


def serialize_report_transactions(transactions):
    return AuctionTransactionDetailReportSerializer(transactions, many=True).data


def store_auction_report_snapshot(data, report_status):
    transactions = data.get("transactions") or data.get("all_transactions") or []
    serialized_transactions = serialize_report_transactions(transactions)
    summary = data.get("summary", {})
    total_amount = (
        summary.get("status_amount")
        or summary.get("total_amount")
        or Decimal("0.00")
    )
    transaction_count = (
        summary.get("status_count")
        if summary.get("status_count") is not None
        else summary.get("total_transactions", 0)
    )

    return AuctionReport.objects.create(
        report_status=report_status,
        transaction_count=transaction_count or 0,
        total_amount=total_amount or Decimal("0.00"),
        summary=make_json_ready(summary),
        payment_status_breakdown=make_json_ready(data.get("payment_status_breakdown", [])),
        transactions=make_json_ready(serialized_transactions),
    )


class AuctionTransactionReportView(View):
    """Generate comprehensive auction transaction reports."""

    def get_search_query(self):
        return (self.request.GET.get("search") or "").strip()

    def apply_search_filter(self, queryset):
        query = self.get_search_query()
        if not query:
            return queryset
        return queryset.filter(
            Q(member__name__icontains=query)
            | Q(member__primary_phone__icontains=query)
            | Q(member__secondary_phone__icontains=query)
            | Q(relative__non_member_id__icontains=query)
            | Q(relative__name__icontains=query)
            | Q(relative__phone_1__icontains=query)
            | Q(relative__phone_2__icontains=query)
        )

    def get_report_data(self):
        """Calculate report data from database."""
        
        # Get all transactions
        all_transactions = self.apply_search_filter(
            AuctionTransaction.objects.select_related("member", "relative", "item", "receipt").all()
        )

        # Calculate totals
        total_count = all_transactions.count()
        
        totals = all_transactions.aggregate(
            total_amount=Sum("price"),
            paid_amount=Sum("price", filter=Q(payment_status="Paid")),
            unpaid_amount=Sum("price", filter=Q(payment_status="Unpaid")),
            paid_count=Count("id", filter=Q(payment_status="Paid")),
            unpaid_count=Count("id", filter=Q(payment_status="Unpaid")),
        )

        total_amount = totals["total_amount"] or Decimal("0.00")
        paid_amount = totals["paid_amount"] or Decimal("0.00")
        unpaid_amount = totals["unpaid_amount"] or Decimal("0.00")
        paid_count = totals["paid_count"] or 0
        unpaid_count = totals["unpaid_count"] or 0

        # Calculate percentages
        paid_percentage = (
            (paid_count / total_count * 100) if total_count > 0 else 0
        )
        unpaid_percentage = (
            (unpaid_count / total_count * 100) if total_count > 0 else 0
        )

        average_amount = (
            (total_amount / total_count) if total_count > 0 else Decimal("0.00")
        )

        # Payment status breakdown
        payment_breakdown = all_transactions.values("payment_status").annotate(
            count=Count("id"),
            total_amount=Sum("price")
        ).order_by("payment_status")

        breakdown_data = []
        for item in payment_breakdown:
            status_val = item["payment_status"]
            status_count = item["count"]
            status_total = item["total_amount"] or Decimal("0.00")
            status_percentage = (status_count / total_count * 100) if total_count > 0 else 0
            
            breakdown_data.append({
                "payment_status": status_val,
                "count": status_count,
                "total_amount": status_total,
                "percentage": round(status_percentage, 2),
            })

        # Summary data
        summary_data = {
            "total_transactions": total_count,
            "total_amount": total_amount,
            "total_paid_amount": paid_amount,
            "total_unpaid_amount": unpaid_amount,
            "paid_count": paid_count,
            "unpaid_count": unpaid_count,
            "paid_percentage": round(paid_percentage, 2),
            "unpaid_percentage": round(unpaid_percentage, 2),
            "average_transaction_amount": round(average_amount, 2),
            "report_date": datetime.now(),
        }

        return {
            "summary": summary_data,
            "payment_status_breakdown": breakdown_data,
            "transactions": all_transactions,
            "all_transactions": all_transactions,
        }

    def get(self, request, *args, **kwargs):
        """Return report in requested format."""
        
        format_type = request.GET.get("format", "json").lower()
        
        if format_type == "json":
            return self.get_json_report()
        elif format_type == "csv":
            return self.get_csv_report()
        elif format_type == "pdf":
            return self.get_pdf_report()
        elif format_type == "excel":
            return self.get_excel_report()
        else:
            return HttpResponse(
                json.dumps({"error": "Invalid format. Supported formats: json, csv, pdf, excel"}),
                content_type="application/json",
                status=400,
            )

    def get_json_report(self):
        """Generate JSON format report."""
        
        data = self.get_report_data()
        store_auction_report_snapshot(data, AuctionReport.ReportStatus.ALL)
        limit, offset = get_limit_offset(self.request)
        transaction_count = data["transactions"].count()
        paged_transactions = data["transactions"][offset : offset + limit]
        payload = {
            "summary": data["summary"],
            "payment_status_breakdown": data["payment_status_breakdown"],
            "transactions": AuctionTransactionDetailReportSerializer(paged_transactions, many=True).data,
            "pagination": {
                "count": transaction_count,
                "limit": limit,
                "offset": offset,
            },
        }
        
        return HttpResponse(
            json.dumps(payload, default=str),
            content_type="application/json",
            status=200,
        )

    def get_csv_report(self):
        """Generate CSV format report."""
        
        data = self.get_report_data()
        store_auction_report_snapshot(data, AuctionReport.ReportStatus.ALL)
        summary = data["summary"]
        transactions = data["all_transactions"]

        output = StringIO()
        writer = csv.writer(output)

        # Write summary section
        writer.writerow(["AUCTION TRANSACTION REPORT"])
        writer.writerow(["Generated Date", summary["report_date"]])
        writer.writerow([])

        # Summary statistics
        writer.writerow(["SUMMARY"])
        writer.writerow(["Total Transactions", summary["total_transactions"]])
        writer.writerow(["Total Amount", f"Rs. {summary['total_amount']}"])
        writer.writerow(["Paid Amount", f"Rs. {summary['total_paid_amount']}"])
        writer.writerow(["Unpaid Amount", f"Rs. {summary['total_unpaid_amount']}"])
        writer.writerow(["Paid Count", summary["paid_count"]])
        writer.writerow(["Unpaid Count", summary["unpaid_count"]])
        writer.writerow(["Paid Percentage", f"{summary['paid_percentage']}%"])
        writer.writerow(["Unpaid Percentage", f"{summary['unpaid_percentage']}%"])
        writer.writerow(["Average Transaction Amount", f"Rs. {summary['average_transaction_amount']}"])
        writer.writerow([])

        # Payment status breakdown
        writer.writerow(["PAYMENT STATUS BREAKDOWN"])
        writer.writerow(["Status", "Count", "Total Amount", "Percentage"])
        for item in data["payment_status_breakdown"]:
            writer.writerow([
                item["payment_status"],
                item["count"],
                f"Rs. {item['total_amount']}",
                f"{item['percentage']}%",
            ])
        writer.writerow([])

        # Detailed transactions
        writer.writerow(["DETAILED TRANSACTIONS"])
        writer.writerow([
            "ID", "Member/Non-Member ID", "Member/Non-Member Name", "Source Type", "Item",
            "Token Number", "Price", "Payment Status", "Payment Date"
        ])
        
        for txn in transactions:
            writer.writerow([
                txn.id,
                get_transaction_source_id(txn),
                get_transaction_source_name(txn),
                get_transaction_source_type(txn),
                txn.item.auction_item_name,
                txn.token_number,
                f"Rs. {txn.price}",
                txn.payment_status,
                get_transaction_payment_date(txn),
            ])

        # Generate HTTP response
        response = HttpResponse(output.getvalue(), content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="auction_transaction_report.csv"'
        return response

    def get_excel_report(self):
        """Generate Excel format report."""
        
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return HttpResponse(
                json.dumps({"error": "openpyxl not installed. Install with: pip install openpyxl"}),
                content_type="application/json",
                status=400,
            )

        data = self.get_report_data()
        store_auction_report_snapshot(data, AuctionReport.ReportStatus.ALL)
        summary = data["summary"]
        transactions = data["all_transactions"]

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        # Styling
        header_fill = PatternFill(start_color="8B5E34", end_color="8B5E34", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        title_font = Font(size=14, bold=True)
        subheader_font = Font(bold=True, size=11)
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        row = 1

        # Title
        worksheet.merge_cells(f"A{row}:J{row}")
        title_cell = worksheet[f"A{row}"]
        title_cell.value = "AUCTION TRANSACTION REPORT"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center")
        row += 1

        # Report date
        worksheet[f"A{row}"] = "Generated Date:"
        worksheet[f"B{row}"] = summary["report_date"].strftime("%Y-%m-%d %H:%M:%S")
        row += 2

        # Summary section
        worksheet[f"A{row}"] = "SUMMARY"
        worksheet[f"A{row}"].font = subheader_font
        row += 1

        summary_rows = [
            ("Total Transactions", summary["total_transactions"]),
            ("Total Amount (Rs.)", summary["total_amount"]),
            ("Paid Amount (Rs.)", summary["total_paid_amount"]),
            ("Unpaid Amount (Rs.)", summary["total_unpaid_amount"]),
            ("Paid Count", summary["paid_count"]),
            ("Unpaid Count", summary["unpaid_count"]),
            ("Paid Percentage (%)", summary["paid_percentage"]),
            ("Unpaid Percentage (%)", summary["unpaid_percentage"]),
            ("Average Transaction Amount (Rs.)", summary["average_transaction_amount"]),
        ]

        for label, value in summary_rows:
            worksheet[f"A{row}"] = label
            worksheet[f"A{row}"].font = Font(bold=True)
            worksheet[f"B{row}"] = value
            row += 1

        row += 1

        # Payment status breakdown
        worksheet[f"A{row}"] = "PAYMENT STATUS BREAKDOWN"
        worksheet[f"A{row}"].font = subheader_font
        row += 1

        headers = ["Status", "Count", "Total Amount (Rs.)", "Percentage (%)"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border

        row += 1

        for item in data["payment_status_breakdown"]:
            worksheet[f"A{row}"] = item["payment_status"]
            worksheet[f"B{row}"] = item["count"]
            worksheet[f"C{row}"] = item["total_amount"]
            worksheet[f"D{row}"] = item["percentage"]
            for col in range(1, 5):
                worksheet.cell(row=row, column=col).border = border
            row += 1

        row += 1

        # Detailed transactions
        worksheet[f"A{row}"] = "DETAILED TRANSACTIONS"
        worksheet[f"A{row}"].font = subheader_font
        row += 1

        headers = [
            "ID", "Member/Non-Member ID", "Member/Non-Member", "Source Type", "Item",
            "Token", "Price (Rs.)", "Payment Status", "Payment Date"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border

        row += 1

        for txn in transactions:
            worksheet[f"A{row}"] = txn.id
            worksheet[f"B{row}"] = get_transaction_source_id(txn)
            worksheet[f"C{row}"] = get_transaction_source_name(txn)
            worksheet[f"D{row}"] = get_transaction_source_type(txn)
            worksheet[f"E{row}"] = txn.item.auction_item_name
            worksheet[f"F{row}"] = txn.token_number
            worksheet[f"G{row}"] = txn.price
            worksheet[f"H{row}"] = txn.payment_status
            worksheet[f"I{row}"] = get_transaction_payment_date(txn)

            for col in range(1, 10):
                worksheet.cell(row=row, column=col).border = border

            row += 1

        # Set column widths
        worksheet.column_dimensions["A"].width = 12
        worksheet.column_dimensions["B"].width = 20
        worksheet.column_dimensions["C"].width = 20
        worksheet.column_dimensions["D"].width = 20
        worksheet.column_dimensions["E"].width = 10
        worksheet.column_dimensions["F"].width = 12
        worksheet.column_dimensions["G"].width = 15
        worksheet.column_dimensions["H"].width = 18
        worksheet.column_dimensions["I"].width = 18

        # Generate response
        return build_excel_response(workbook, "auction_transaction_report.xlsx")

    def get_pdf_report(self):
        """Generate PDF format report using fpdf2."""
        
        try:
            from fpdf import FPDF
        except ImportError:
            return HttpResponse(
                json.dumps({"error": "fpdf2 not installed. Install with: pip install fpdf2"}),
                content_type="application/json",
                status=400,
            )

        data = self.get_report_data()
        store_auction_report_snapshot(data, AuctionReport.ReportStatus.ALL)
        summary = data["summary"]
        transactions = data["all_transactions"]

        pdf = FPDF(orientation="L")
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        # Title
        pdf.set_font("Arial", "B", size=16)
        pdf.cell(0, 10, "AUCTION TRANSACTION REPORT", ln=True, align="C")
        pdf.ln(5)

        # Report date
        pdf.set_font("Arial", size=10)
        pdf.cell(0, 10, f"Generated: {summary['report_date'].strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        pdf.ln(5)

        # Summary section
        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 10, "SUMMARY", ln=True)
        pdf.set_font("Arial", size=10)

        summary_items = [
            f"Total Transactions: {summary['total_transactions']}",
            f"Total Amount: Rs. {summary['total_amount']}",
            f"Paid Amount: Rs. {summary['total_paid_amount']}",
            f"Unpaid Amount: Rs. {summary['total_unpaid_amount']}",
            f"Paid Count: {summary['paid_count']} ({summary['paid_percentage']}%)",
            f"Unpaid Count: {summary['unpaid_count']} ({summary['unpaid_percentage']}%)",
            f"Average Amount: Rs. {summary['average_transaction_amount']}",
        ]

        for item in summary_items:
            pdf.cell(0, 8, item, ln=True)

        pdf.ln(5)

        # Payment status breakdown
        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 10, "PAYMENT STATUS BREAKDOWN", ln=True)
        pdf.set_font("Arial", size=10)

        for item in data["payment_status_breakdown"]:
            pdf.cell(0, 8, f"{item['payment_status']}: {item['count']} transactions (Rs. {item['total_amount']}) - {item['percentage']}%", ln=True)

        pdf.ln(5)

        # Detailed transactions table
        pdf.set_font("Arial", "B", size=11)
        pdf.cell(0, 10, "DETAILED TRANSACTIONS", ln=True)

        columns = [
            ("ID", 10, 8),
            ("Source ID", 22, 14),
            ("Name", 48, 20),
            ("Type", 18, 12),
            ("Item", 44, 20),
            ("Token", 14, 8),
            ("Price", 24, 14),
            ("Status", 20, 10),
            ("Receipt", 24, 14),
            ("Payment Date", 30, 16),
        ]
        draw_pdf_table_header(pdf, columns)

        for txn in transactions:
            draw_pdf_table_row(
                pdf,
                columns,
                [
                    txn.id,
                    get_transaction_source_id(txn),
                    get_transaction_source_name(txn),
                    get_transaction_source_type(txn),
                    txn.item.auction_item_name,
                    txn.token_number,
                    f"Rs. {txn.price}",
                    txn.payment_status,
                    get_transaction_receipt_no(txn),
                    get_transaction_payment_date(txn),
                ],
            )

        # Generate response
        pdf_bytes = normalize_pdf_output(pdf.output())
        
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="auction_transaction_report.pdf"'
        return response


class AuctionPaymentStatusReportView(View):
    """
    Report for either Paid or Unpaid auction transactions.

    Subclasses must set `payment_status_value` to one of:
      - "Paid"
      - "Unpaid"
    """

    payment_status_value: str | None = None

    def get_payment_status_label(self) -> str:
        if self.payment_status_value == "Paid":
            return "PAID"
        if self.payment_status_value == "Unpaid":
            return "UNPAID"
        return str(self.payment_status_value or "").upper()

    def get_search_query(self):
        return (self.request.GET.get("search") or "").strip()

    def apply_search_filter(self, queryset):
        query = self.get_search_query()
        if not query:
            return queryset
        return queryset.filter(
            Q(member__name__icontains=query)
            | Q(member__primary_phone__icontains=query)
            | Q(member__secondary_phone__icontains=query)
            | Q(relative__non_member_id__icontains=query)
            | Q(relative__name__icontains=query)
            | Q(relative__phone_1__icontains=query)
            | Q(relative__phone_2__icontains=query)
        )

    def get_all_transactions(self):
        queryset = AuctionTransaction.objects.select_related("member", "relative", "item", "receipt")
        return self.apply_search_filter(queryset)

    def get_filtered_transactions(self):
        return self.get_all_transactions().filter(payment_status=self.payment_status_value).order_by("-created_at")

    def get_report_data(self):
        if not self.payment_status_value:
            return {
                "summary": {},
                "transactions": [],
                "payment_status_breakdown": [],
            }

        all_transactions = self.get_all_transactions()
        total_count = all_transactions.count()

        totals = all_transactions.aggregate(
            total_amount=Sum("price"),
            paid_amount=Sum("price", filter=Q(payment_status="Paid")),
            unpaid_amount=Sum("price", filter=Q(payment_status="Unpaid")),
            paid_count=Count("id", filter=Q(payment_status="Paid")),
            unpaid_count=Count("id", filter=Q(payment_status="Unpaid")),
        )

        total_amount = totals["total_amount"] or Decimal("0.00")
        paid_amount = totals["paid_amount"] or Decimal("0.00")
        unpaid_amount = totals["unpaid_amount"] or Decimal("0.00")
        paid_count = totals["paid_count"] or 0
        unpaid_count = totals["unpaid_count"] or 0

        paid_percentage = (paid_count / total_count * 100) if total_count > 0 else 0
        unpaid_percentage = (unpaid_count / total_count * 100) if total_count > 0 else 0

        status_transactions = self.get_filtered_transactions()
        status_count = status_transactions.count()
        status_amount = status_transactions.aggregate(total_amount=Sum("price"))["total_amount"] or Decimal("0.00")

        other_count = unpaid_count if self.payment_status_value == "Paid" else paid_count
        other_amount = unpaid_amount if self.payment_status_value == "Paid" else paid_amount
        other_label = "Unpaid" if self.payment_status_value == "Paid" else "Paid"

        payment_breakdown_qs = all_transactions.values("payment_status").annotate(
            count=Count("id"),
            total_amount=Sum("price"),
        ).order_by("payment_status")

        payment_status_breakdown = []
        for item in payment_breakdown_qs:
            status_val = item["payment_status"]
            status_count_b = item["count"]
            status_total = item["total_amount"] or Decimal("0.00")
            status_percentage = (status_count_b / total_count * 100) if total_count > 0 else 0
            payment_status_breakdown.append(
                {
                    "payment_status": status_val,
                    "count": status_count_b,
                    "total_amount": status_total,
                    "percentage": round(status_percentage, 2),
                }
            )

        summary_data = {
            "report_date": datetime.now(),
            "total_transactions": total_count,
            "total_amount": total_amount,
            "paid_amount": paid_amount,
            "unpaid_amount": unpaid_amount,
            "paid_count": paid_count,
            "unpaid_count": unpaid_count,
            "paid_percentage": round(paid_percentage, 2),
            "unpaid_percentage": round(unpaid_percentage, 2),
            "status": self.payment_status_value,
            "status_count": status_count,
            "status_amount": status_amount,
            "other_label": other_label,
            "other_count": other_count,
            "other_amount": other_amount,
            "average_transaction_amount": round((total_amount / total_count), 2) if total_count > 0 else Decimal("0.00"),
        }

        return {
            "summary": summary_data,
            "payment_status_breakdown": payment_status_breakdown,
            "transactions": status_transactions,
        }

    def get(self, request, *args, **kwargs):
        format_type = request.GET.get("format", "json").lower()

        if format_type == "json":
            return self.get_json_report()
        if format_type == "csv":
            return self.get_csv_report()
        if format_type == "excel":
            return self.get_excel_report()
        if format_type == "pdf":
            return self.get_pdf_report()

        return HttpResponse(
            json.dumps({"error": "Invalid format. Supported formats: json, csv, excel, pdf"}),
            content_type="application/json",
            status=400,
        )

    def get_json_report(self):
        data = self.get_report_data()
        store_auction_report_snapshot(data, self.payment_status_value)
        limit, offset = get_limit_offset(self.request)
        transaction_count = data["transactions"].count()
        paged_transactions = data["transactions"][offset : offset + limit]
        serialized_transactions = AuctionTransactionDetailReportSerializer(paged_transactions, many=True).data
        payload = {
            "summary": data["summary"],
            "payment_status_breakdown": data["payment_status_breakdown"],
            "transactions": serialized_transactions,
            "pagination": {
                "count": transaction_count,
                "limit": limit,
                "offset": offset,
            },
        }
        return HttpResponse(
            json.dumps(payload, default=str),
            content_type="application/json",
            status=200,
        )

    def get_csv_report(self):
        data = self.get_report_data()
        store_auction_report_snapshot(data, self.payment_status_value)
        summary = data["summary"]
        transactions = data["transactions"]

        output = StringIO()
        writer = csv.writer(output)

        status_label = self.get_payment_status_label()

        writer.writerow([f"{status_label} AUCTION TRANSACTION REPORT"])
        writer.writerow(["Generated Date", summary["report_date"]])
        writer.writerow([])

        writer.writerow(["SUMMARY"])
        writer.writerow(["Total Transactions (All)", summary["total_transactions"]])
        writer.writerow(["Total Amount (All)", f"Rs. {summary['total_amount']}"])
        writer.writerow([f"{status_label} Count", summary["status_count"]])
        writer.writerow([f"{status_label} Amount", f"Rs. {summary['status_amount']}"])
        writer.writerow([f"Other ({summary['other_label']}) Count", summary["other_count"]])
        writer.writerow([f"Other ({summary['other_label']}) Amount", f"Rs. {summary['other_amount']}"])
        writer.writerow(["Paid Percentage", f"{summary['paid_percentage']}%"])
        writer.writerow(["Unpaid Percentage", f"{summary['unpaid_percentage']}%"])
        writer.writerow([])

        writer.writerow(["PAYMENT STATUS BREAKDOWN"])
        writer.writerow(["Status", "Count", "Total Amount", "Percentage"])
        for item in data["payment_status_breakdown"]:
            writer.writerow(
                [
                    item["payment_status"],
                    item["count"],
                    f"Rs. {item['total_amount']}",
                    f"{item['percentage']}%",
                ]
            )
        writer.writerow([])

        writer.writerow(["DETAILED TRANSACTIONS"])
        writer.writerow(
            [
                "ID",
                "Member/Non-Member ID",
                "Member/Non-Member Name",
                "Source Type",
                "Item",
                "Token Number",
                "Price",
                "Payment Status",
                "Payment Date",
            ]
        )

        for txn in transactions:
            writer.writerow(
                [
                    txn.id,
                    get_transaction_source_id(txn),
                    get_transaction_source_name(txn),
                    get_transaction_source_type(txn),
                    txn.item.auction_item_name,
                    txn.token_number,
                    f"Rs. {txn.price}",
                    txn.payment_status,
                    get_transaction_payment_date(txn),
                ]
            )

        response = HttpResponse(output.getvalue(), content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{status_label.lower()}_auction_transaction_report.csv"'
        return response

    def get_excel_report(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            return HttpResponse(
                json.dumps({"error": "openpyxl not installed. Install with: pip install openpyxl"}),
                content_type="application/json",
                status=400,
            )

        data = self.get_report_data()
        store_auction_report_snapshot(data, self.payment_status_value)
        status_label = self.get_payment_status_label()
        workbook = create_payment_status_excel_workbook(data, status_label)
        filename = f"{status_label.lower()}_auction_transaction_report.xlsx"
        return build_excel_response(workbook, filename)

    def get_pdf_report(self):
        try:
            from fpdf import FPDF
        except ImportError:
            return HttpResponse(
                json.dumps({"error": "fpdf2 not installed. Install with: pip install fpdf2"}),
                content_type="application/json",
                status=400,
            )

        data = self.get_report_data()
        store_auction_report_snapshot(data, self.payment_status_value)
        summary = data["summary"]
        transactions = data["transactions"]

        status_label = self.get_payment_status_label()

        pdf = FPDF(orientation="L")
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        pdf.set_font("Arial", "B", size=16)
        pdf.cell(0, 10, f"{status_label} AUCTION TRANSACTION REPORT", ln=True, align="C")
        pdf.ln(4)

        pdf.set_font("Arial", size=10)
        pdf.cell(0, 8, f"Generated: {summary['report_date'].strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        pdf.ln(4)

        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 8, "SUMMARY", ln=True)

        summary_columns = [
            ("Status", 42, 18),
            ("Count", 28, 10),
            ("Amount", 42, 18),
            ("Total Count", 34, 12),
            ("Total Amount", 42, 18),
            ("Paid %", 28, 10),
            ("Unpaid %", 28, 10),
        ]
        draw_pdf_table_header(pdf, summary_columns)
        draw_pdf_table_row(
            pdf,
            summary_columns,
            [
                summary["status"],
                summary["status_count"],
                f"Rs. {summary['status_amount']}",
                summary["total_transactions"],
                f"Rs. {summary['total_amount']}",
                f"{summary['paid_percentage']}%",
                f"{summary['unpaid_percentage']}%",
            ],
        )

        pdf.ln(3)

        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 8, f"{status_label} TRANSACTION DETAILS", ln=True)

        detail_columns = [
            ("ID", 10, 8),
            ("Source ID", 22, 14),
            ("Name", 48, 20),
            ("Type", 18, 12),
            ("Item", 44, 20),
            ("Token", 14, 8),
            ("Price", 24, 14),
            ("Status", 20, 10),
            ("Receipt", 24, 14),
            ("Payment Date", 30, 16),
        ]
        draw_pdf_table_header(pdf, detail_columns)

        for txn in transactions:
            draw_pdf_table_row(
                pdf,
                detail_columns,
                [
                    txn.id,
                    get_transaction_source_id(txn),
                    get_transaction_source_name(txn),
                    get_transaction_source_type(txn),
                    txn.item.auction_item_name,
                    txn.token_number,
                    f"Rs. {txn.price}",
                    txn.payment_status,
                    get_transaction_receipt_no(txn),
                    get_transaction_payment_date(txn),
                ],
            )

        pdf_bytes = normalize_pdf_output(pdf.output(dest="S"))

        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{status_label.lower()}_auction_transaction_report.pdf"'
        return response


class PaidAuctionTransactionReportView(AuctionPaymentStatusReportView):
    payment_status_value = "Paid"


class UnpaidAuctionTransactionReportView(AuctionPaymentStatusReportView):
    payment_status_value = "Unpaid"


class DonationReportView(View):
    """Generate donation reports for member and non-member donors."""

    def get_report_data(self):
        search = self.request.GET.get("search", "").strip()
        donations = get_filtered_donations(search)
        aggregates = donations.aggregate(total_amount=Sum("amount"), donor_count=Count("id"))
        member_aggregates = donations.filter(donor_type=Donation.DonorType.MEMBER).aggregate(
            total_amount=Sum("amount"),
            donor_count=Count("id"),
        )
        non_member_aggregates = donations.filter(donor_type=Donation.DonorType.NON_MEMBER).aggregate(
            total_amount=Sum("amount"),
            donor_count=Count("id"),
        )
        return {
            "search": search,
            "generated_at": datetime.now(),
            "donations": donations,
            "summary": {
                "donor_count": aggregates["donor_count"] or 0,
                "total_amount": aggregates["total_amount"] or Decimal("0.00"),
                "member_count": member_aggregates["donor_count"] or 0,
                "member_amount": member_aggregates["total_amount"] or Decimal("0.00"),
                "non_member_count": non_member_aggregates["donor_count"] or 0,
                "non_member_amount": non_member_aggregates["total_amount"] or Decimal("0.00"),
            },
        }

    def get(self, request, *args, **kwargs):
        format_type = request.GET.get("format", "pdf").lower()
        if format_type != "pdf":
            return HttpResponse(
                json.dumps({"error": "Invalid format. Supported format: pdf"}),
                content_type="application/json",
                status=400,
            )
        return self.get_pdf_report()

    def get_pdf_report(self):
        try:
            from fpdf import FPDF
        except ImportError:
            return HttpResponse(
                json.dumps({"error": "fpdf2 not installed. Install with: pip install fpdf2"}),
                content_type="application/json",
                status=400,
            )

        data = self.get_report_data()
        summary = data["summary"]
        donations = data["donations"]

        pdf = FPDF(orientation="L")
        pdf.add_page()
        pdf.set_font("Arial", "B", size=16)
        pdf.cell(0, 10, "DONATION REPORT", ln=True, align="C")
        pdf.ln(4)

        pdf.set_font("Arial", size=10)
        pdf.cell(0, 8, f"Generated: {data['generated_at'].strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        if data["search"]:
            pdf.cell(0, 8, f"Search: {data['search']}", ln=True)
        pdf.ln(4)

        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 8, "SUMMARY", ln=True)
        summary_columns = [
            ("Total Donors", 38, 16),
            ("Total Amount", 42, 16),
            ("Member Donors", 38, 16),
            ("Member Amount", 42, 16),
            ("Non-Member Donors", 46, 18),
            ("Non-Member Amount", 50, 18),
        ]
        draw_pdf_table_header(pdf, summary_columns)
        draw_pdf_table_row(
            pdf,
            summary_columns,
            [
                summary["donor_count"],
                f"Rs. {summary['total_amount']}",
                summary["member_count"],
                f"Rs. {summary['member_amount']}",
                summary["non_member_count"],
                f"Rs. {summary['non_member_amount']}",
            ],
        )
        pdf.ln(5)

        pdf.set_font("Arial", "B", size=12)
        pdf.cell(0, 8, "DONATION DETAILS", ln=True)
        detail_columns = [
            ("ID", 12, 8),
            ("Donor Type", 28, 14),
            ("Member/Non-member ID", 42, 18),
            ("Donor Name", 62, 26),
            ("Phone", 32, 12),
            ("Amount", 34, 14),
            ("Created", 38, 18),
        ]
        draw_pdf_table_header(pdf, detail_columns)

        for donation in donations:
            draw_pdf_table_row(
                pdf,
                detail_columns,
                [
                    donation.id,
                    donation.donor_type,
                    get_donation_source_id(donation),
                    donation.donor_name,
                    donation.phone,
                    f"Rs. {donation.amount}",
                    donation.created_at.strftime("%Y-%m-%d"),
                ],
            )

        pdf_bytes = normalize_pdf_output(pdf.output(dest="S"))
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'inline; filename="donation_report.pdf"'
        return response
